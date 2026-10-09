import csv
import re
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook

from config import DE_LCL_XLSX_PATH
from db import get_connection
from util import now_stamp

HEADER_MAP = {
    "consignee": "consignee",
    "consignee name": "consigneeName",
    "orga code": "orgaCode",
    "orgacode": "orgaCode",
    "remark": "remark",
    "senator": "senator",
    "delivery agent": "deliveryAgent",
    "deliveryagent": "deliveryAgent",
}

LIST_COLS = "ID, Consignee, ConsigneeName, OrgaCode, Remark, Senator, DeliveryAgent"
LIST_ORDER = "Senator, ConsigneeName, OrgaCode, ID"


def normalize_header(value):
    text = str(value or "").strip().lower().replace("\ufeff", "")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def clean_text(value):
    text = str(value or "").replace("\u00a0", " ").replace("\ufeff", "")
    return " ".join(text.replace("\r\n", "\n").replace("\r", "\n").split())


def decode_csv(data):
    if not data:
        return ""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def search_text(fields):
    return " ".join(
        part
        for part in (
            fields["consignee"],
            fields["consigneeName"],
            fields["orgaCode"],
            fields["remark"],
            fields["senator"],
            fields["deliveryAgent"],
        )
        if part
    )


def row_to_dict(row):
    return {
        "id": row["ID"],
        "consignee": row["Consignee"],
        "consigneeName": row["ConsigneeName"],
        "orgaCode": row["OrgaCode"],
        "remark": row["Remark"],
        "senator": row["Senator"],
        "deliveryAgent": row["DeliveryAgent"],
    }


def map_headers(raw_header):
    mapping = {}
    for index, cell in enumerate(raw_header):
        key = HEADER_MAP.get(normalize_header(cell))
        if key and key not in mapping:
            mapping[key] = index
    return mapping


def records_from_rows(mapping, rows):
    if "consignee" not in mapping and "consigneeName" not in mapping and "orgaCode" not in mapping:
        return [], "File must include consignee, consignee name, or orga code."
    records = []
    for raw in rows:
        if not raw or not any(cell_text(cell) for cell in raw):
            continue

        def col(name):
            index = mapping.get(name)
            if index is None or index >= len(raw):
                return ""
            return cell_text(raw[index])

        fields = {
            "consignee": col("consignee"),
            "consigneeName": col("consigneeName"),
            "orgaCode": col("orgaCode"),
            "remark": col("remark"),
            "senator": col("senator"),
            "deliveryAgent": col("deliveryAgent"),
        }
        if not any(fields[key] for key in ("consignee", "consigneeName", "orgaCode")):
            continue
        fields["searchText"] = search_text(fields)
        records.append(fields)
    if not records:
        return [], "File contains no DE-LCL rows."
    return records, None


def parse_delcl_csv(data):
    text = decode_csv(data)
    if not text.strip():
        return [], "CSV is empty."
    reader = csv.reader(StringIO(text))
    try:
        raw_header = next(reader)
    except StopIteration:
        return [], "CSV header row is missing."
    return records_from_rows(map_headers(raw_header), reader)


def parse_delcl_xlsx(data):
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return [], "Excel file could not be read."
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_header = next(rows)
        except StopIteration:
            return [], "Excel header row is missing."
        return records_from_rows(map_headers(raw_header), rows)
    finally:
        workbook.close()


def parse_delcl_file(filename, data):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return parse_delcl_csv(data)
    return parse_delcl_xlsx(data)


def import_delcl_file(filename, data):
    records, error = parse_delcl_file(filename, data)
    if error:
        return None, error
    stamp = now_stamp()
    stored_name = (filename or "DE-LCL.xlsx")[:200]
    with get_connection() as conn:
        conn.execute("BEGIN IMMEDIATE")
        conn.execute("DELETE FROM DeLclRecords")
        conn.executemany(
            """
            INSERT INTO DeLclRecords (
                Consignee, ConsigneeName, OrgaCode, Remark, Senator, DeliveryAgent, SearchText
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    item["consignee"],
                    item["consigneeName"],
                    item["orgaCode"],
                    item["remark"],
                    item["senator"],
                    item["deliveryAgent"],
                    item["searchText"],
                )
                for item in records
            ],
        )
        conn.execute(
            """
            INSERT INTO DeLclImportMeta (ID, Filename, ImportedAt, RowCount)
            VALUES (1, ?, ?, ?)
            ON CONFLICT(ID) DO UPDATE SET
                Filename=excluded.Filename,
                ImportedAt=excluded.ImportedAt,
                RowCount=excluded.RowCount
            """,
            (stored_name, stamp, len(records)),
        )
        conn.execute("ANALYZE DeLclRecords")
    return {"filename": stored_name, "rowCount": len(records), "importedAt": stamp}, None


def ensure_delcl_loaded():
    with get_connection() as conn:
        count = conn.execute("SELECT COUNT(*) FROM DeLclRecords").fetchone()[0]
    if count:
        return None, None
    path = Path(DE_LCL_XLSX_PATH)
    if not path.is_file():
        return None, None
    return import_delcl_file("DE-LCL.xlsx", path.read_bytes())


def list_delcl(query="", page=1, page_size=50):
    q = clean_text(query)
    page = max(int(page or 1), 1)
    page_size = min(max(int(page_size or 50), 1), 200)
    offset = (page - 1) * page_size
    with get_connection() as conn:
        meta_row = conn.execute("SELECT * FROM DeLclImportMeta WHERE ID=1").fetchone()
        meta_total = int(meta_row["RowCount"]) if meta_row else 0
        if not q:
            total = meta_total or conn.execute("SELECT COUNT(*) FROM DeLclRecords").fetchone()[0]
            rows = conn.execute(
                f"""
                SELECT {LIST_COLS} FROM DeLclRecords
                ORDER BY {LIST_ORDER}
                LIMIT ? OFFSET ?
                """,
                (page_size, offset),
            ).fetchall()
        else:
            like = f"%{q}%"
            where = "WHERE SearchText LIKE ? COLLATE NOCASE"
            params = [like]
            total = conn.execute(f"SELECT COUNT(*) FROM DeLclRecords {where}", params).fetchone()[0]
            rows = conn.execute(
                f"""
                SELECT {LIST_COLS} FROM DeLclRecords {where}
                ORDER BY {LIST_ORDER}
                LIMIT ? OFFSET ?
                """,
                params + [page_size, offset],
            ).fetchall()
    meta = {
        "filename": meta_row["Filename"] if meta_row else "",
        "importedAt": meta_row["ImportedAt"] if meta_row else "",
        "rowCount": meta_row["RowCount"] if meta_row else total,
    }
    return {
        "data": [row_to_dict(row) for row in rows],
        "pagination": {
            "page": page,
            "pageSize": page_size,
            "total": total,
            "totalPages": max((total + page_size - 1) // page_size, 1),
        },
        "meta": meta,
    }


def bump_delcl_row_count(conn, delta=1):
    delta = int(delta)
    conn.execute(
        """
        INSERT INTO DeLclImportMeta (ID, Filename, ImportedAt, RowCount)
        VALUES (1, 'manual', '', MAX(?, 0))
        ON CONFLICT(ID) DO UPDATE SET RowCount=MAX(DeLclImportMeta.RowCount + ?, 0)
        """,
        (delta, delta),
    )


def parse_delcl_payload(payload):
    fields = {
        "consignee": clean_text((payload or {}).get("consignee")),
        "consigneeName": clean_text((payload or {}).get("consigneeName")),
        "orgaCode": clean_text((payload or {}).get("orgaCode")),
        "remark": clean_text((payload or {}).get("remark")),
        "senator": clean_text((payload or {}).get("senator")),
        "deliveryAgent": clean_text((payload or {}).get("deliveryAgent")),
    }
    if not any(fields[key] for key in ("consignee", "consigneeName", "orgaCode")):
        return None, "Consignee, consignee name, or orga code is required."
    fields["searchText"] = search_text(fields)
    return fields, None


def create_delcl(payload):
    fields, error = parse_delcl_payload(payload)
    if error:
        return None, error
    with get_connection() as conn:
        cur = conn.execute(
            """
            INSERT INTO DeLclRecords (
                Consignee, ConsigneeName, OrgaCode, Remark, Senator, DeliveryAgent, SearchText
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                fields["consignee"],
                fields["consigneeName"],
                fields["orgaCode"],
                fields["remark"],
                fields["senator"],
                fields["deliveryAgent"],
                fields["searchText"],
            ),
        )
        bump_delcl_row_count(conn)
        row = conn.execute("SELECT * FROM DeLclRecords WHERE ID=?", (cur.lastrowid,)).fetchone()
    return row_to_dict(row), None


def update_delcl(record_id, payload):
    fields, error = parse_delcl_payload(payload)
    if error:
        return None, error
    with get_connection() as conn:
        existing = conn.execute("SELECT ID FROM DeLclRecords WHERE ID=?", (record_id,)).fetchone()
        if not existing:
            return None, "Record not found"
        conn.execute(
            """
            UPDATE DeLclRecords
            SET Consignee=?, ConsigneeName=?, OrgaCode=?, Remark=?, Senator=?, DeliveryAgent=?, SearchText=?
            WHERE ID=?
            """,
            (
                fields["consignee"],
                fields["consigneeName"],
                fields["orgaCode"],
                fields["remark"],
                fields["senator"],
                fields["deliveryAgent"],
                fields["searchText"],
                record_id,
            ),
        )
        row = conn.execute("SELECT * FROM DeLclRecords WHERE ID=?", (record_id,)).fetchone()
    return row_to_dict(row), None


def delete_delcl(record_id):
    with get_connection() as conn:
        row = conn.execute("SELECT * FROM DeLclRecords WHERE ID=?", (record_id,)).fetchone()
        if not row:
            return None, "Record not found"
        conn.execute("DELETE FROM DeLclRecords WHERE ID=?", (record_id,))
        bump_delcl_row_count(conn, -1)
    return row_to_dict(row), None
