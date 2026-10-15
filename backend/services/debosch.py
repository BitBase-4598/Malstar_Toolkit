import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from config import BOSCH_DE_XLSX_PATH
from db import get_connection
from util import now_stamp

HEADER_MAP = {
    "full name": "fullName",
    "address 1": "address1",
    "address": "address1",
    "unloco": "unloco",
    "unlocode": "unloco",
    "organization code": "organizationCode",
    "orga code": "organizationCode",
    "org code": "organizationCode",
    "maersk be code": "beCode",
    "be code": "beCode",
    "maersk cmd code": "cmdCode",
    "cmd code": "cmdCode",
    "maersk bp code": "bpCode",
    "bp code": "bpCode",
    "city": "city",
    "postcode": "postcode",
    "post code": "postcode",
}

FIELDS = (
    "fullName",
    "address1",
    "unloco",
    "organizationCode",
    "beCode",
    "cmdCode",
    "bpCode",
    "city",
    "postcode",
)

LIST_COLS = (
    "ID, FullName, Address1, Unloco, OrganizationCode, BeCode, CmdCode, BpCode, City, Postcode"
)
LIST_ORDER = "FullName, OrganizationCode, ID"


def normalize_header(value):
    text = str(value or "").strip().lower().replace("\ufeff", "")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def clean_text(value):
    text = str(value or "").replace("\u00a0", " ").replace("\ufeff", "")
    return " ".join(text.replace("\r\n", "\n").replace("\r", "\n").split())


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def search_text(fields):
    return " ".join(fields[key] for key in FIELDS if fields[key])


def row_to_dict(row):
    return {
        "id": row["ID"],
        "fullName": row["FullName"],
        "address1": row["Address1"],
        "unloco": row["Unloco"],
        "organizationCode": row["OrganizationCode"],
        "beCode": row["BeCode"],
        "cmdCode": row["CmdCode"],
        "bpCode": row["BpCode"],
        "city": row["City"],
        "postcode": row["Postcode"],
    }


def map_headers(raw_header):
    mapping = {}
    for index, cell in enumerate(raw_header):
        key = HEADER_MAP.get(normalize_header(cell))
        if key and key not in mapping:
            mapping[key] = index
    return mapping


def records_from_rows(mapping, rows):
    if "fullName" not in mapping and "organizationCode" not in mapping:
        return [], "File must include Full Name or Organization Code."
    records = []
    for raw in rows:
        if not raw or not any(cell_text(cell) for cell in raw):
            continue

        def col(name):
            index = mapping.get(name)
            if index is None or index >= len(raw):
                return ""
            return cell_text(raw[index])

        fields = {key: col(key) for key in FIELDS}
        if not fields["fullName"] and not fields["organizationCode"]:
            continue
        fields["searchText"] = search_text(fields)
        records.append(fields)
    if not records:
        return [], "File contains no DE-BOSCH rows."
    return records, None


def parse_debosch_xlsx(data):
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return [], "Excel file could not be read."
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_header = next(rows)
        except StopIteration:
            return [], "Excel header row is missing."
        return records_from_rows(map_headers(raw_header), rows)
    finally:
        workbook.close()


def import_debosch_file(filename, data):
    records, error = parse_debosch_xlsx(data)
    if error:
        return None, error
    stamp = now_stamp()
    stored_name = (filename or "Bosch-DE.xlsx")[:200]
    with get_connection() as conn:
        conn.execute("BEGIN IMMEDIATE")
        conn.execute("DELETE FROM DeBoschRecords")
        conn.executemany(
            """
            INSERT INTO DeBoschRecords (
                FullName, Address1, Unloco, OrganizationCode, BeCode, CmdCode, BpCode,
                City, Postcode, SearchText
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    item["fullName"],
                    item["address1"],
                    item["unloco"],
                    item["organizationCode"],
                    item["beCode"],
                    item["cmdCode"],
                    item["bpCode"],
                    item["city"],
                    item["postcode"],
                    item["searchText"],
                )
                for item in records
            ],
        )
        conn.execute(
            """
            INSERT INTO DeBoschImportMeta (ID, Filename, ImportedAt, RowCount)
            VALUES (1, ?, ?, ?)
            ON CONFLICT(ID) DO UPDATE SET
                Filename=excluded.Filename,
                ImportedAt=excluded.ImportedAt,
                RowCount=excluded.RowCount
            """,
            (stored_name, stamp, len(records)),
        )
        conn.execute("ANALYZE DeBoschRecords")
    return {"filename": stored_name, "rowCount": len(records), "importedAt": stamp}, None


def ensure_debosch_loaded():
    with get_connection() as conn:
        count = conn.execute("SELECT COUNT(*) FROM DeBoschRecords").fetchone()[0]
    if count:
        return None, None
    path = Path(BOSCH_DE_XLSX_PATH)
    if not path.is_file():
        return None, None
    return import_debosch_file("Bosch-DE.xlsx", path.read_bytes())


def list_debosch(query="", page=1, page_size=50):
    q = clean_text(query)
    page = max(int(page or 1), 1)
    page_size = min(max(int(page_size or 50), 1), 200)
    offset = (page - 1) * page_size
    with get_connection() as conn:
        meta_row = conn.execute("SELECT * FROM DeBoschImportMeta WHERE ID=1").fetchone()
        meta_total = int(meta_row["RowCount"]) if meta_row else 0
        if not q:
            total = meta_total or conn.execute("SELECT COUNT(*) FROM DeBoschRecords").fetchone()[0]
            rows = conn.execute(
                f"""
                SELECT {LIST_COLS} FROM DeBoschRecords
                ORDER BY {LIST_ORDER}
                LIMIT ? OFFSET ?
                """,
                (page_size, offset),
            ).fetchall()
        else:
            like = f"%{q}%"
            where = "WHERE SearchText LIKE ? COLLATE NOCASE"
            params = [like]
            total = conn.execute(f"SELECT COUNT(*) FROM DeBoschRecords {where}", params).fetchone()[0]
            rows = conn.execute(
                f"""
                SELECT {LIST_COLS} FROM DeBoschRecords {where}
                ORDER BY {LIST_ORDER}
                LIMIT ? OFFSET ?
                """,
                params + [page_size, offset],
            ).fetchall()
    meta = {
        "filename": meta_row["Filename"] if meta_row else "",
        "importedAt": meta_row["ImportedAt"] if meta_row else "",
        "rowCount": meta_row["RowCount"] if meta_row else total,
    }
    return {
        "data": [row_to_dict(row) for row in rows],
        "pagination": {
            "page": page,
            "pageSize": page_size,
            "total": total,
            "totalPages": max((total + page_size - 1) // page_size, 1),
        },
        "meta": meta,
    }


def bump_debosch_row_count(conn, delta=1):
    delta = int(delta)
    conn.execute(
        """
        INSERT INTO DeBoschImportMeta (ID, Filename, ImportedAt, RowCount)
        VALUES (1, 'manual', '', MAX(?, 0))
        ON CONFLICT(ID) DO UPDATE SET RowCount=MAX(DeBoschImportMeta.RowCount + ?, 0)
        """,
        (delta, delta),
    )


def parse_debosch_payload(payload):
    source = payload or {}
    fields = {key: clean_text(source.get(key)) for key in FIELDS}
    if not fields["fullName"] and not fields["organizationCode"]:
        return None, "Full name or organization code is required."
    fields["searchText"] = search_text(fields)
    return fields, None


def _insert_values(fields):
    return tuple(fields[key] for key in (*FIELDS, "searchText"))


def create_debosch(payload):
    fields, error = parse_debosch_payload(payload)
    if error:
        return None, error
    with get_connection() as conn:
        cur = conn.execute(
            """
            INSERT INTO DeBoschRecords (
                FullName, Address1, Unloco, OrganizationCode, BeCode, CmdCode, BpCode,
                City, Postcode, SearchText
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            _insert_values(fields),
        )
        bump_debosch_row_count(conn)
        row = conn.execute("SELECT * FROM DeBoschRecords WHERE ID=?", (cur.lastrowid,)).fetchone()
    return row_to_dict(row), None


def update_debosch(record_id, payload):
    fields, error = parse_debosch_payload(payload)
    if error:
        return None, error
    with get_connection() as conn:
        existing = conn.execute("SELECT ID FROM DeBoschRecords WHERE ID=?", (record_id,)).fetchone()
        if not existing:
            return None, "Record not found"
        conn.execute(
            """
            UPDATE DeBoschRecords
            SET FullName=?, Address1=?, Unloco=?, OrganizationCode=?, BeCode=?, CmdCode=?,
                BpCode=?, City=?, Postcode=?, SearchText=?
            WHERE ID=?
            """,
            _insert_values(fields) + (record_id,),
        )
        row = conn.execute("SELECT * FROM DeBoschRecords WHERE ID=?", (record_id,)).fetchone()
    return row_to_dict(row), None


def delete_debosch(record_id):
    with get_connection() as conn:
        row = conn.execute("SELECT * FROM DeBoschRecords WHERE ID=?", (record_id,)).fetchone()
        if not row:
            return None, "Record not found"
        conn.execute("DELETE FROM DeBoschRecords WHERE ID=?", (record_id,))
        bump_debosch_row_count(conn, -1)
    return row_to_dict(row), None
